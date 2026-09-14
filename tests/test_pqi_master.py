"""The PQI master workbook is read by name, validated, and refused when it is unusable."""
import warnings
from collections import Counter

import openpyxl
import pytest

from conftest import ROOT
from pqi_master import PqiMasterError, load_master

MASTER = ROOT / "data" / "RDC_PQI_Master_100_AI_Scoring_v1.0_PILOT.xlsx"
MASTER_SHA256 = "093185fd20a3b29a18c37d03818a546c6fc6c07ffb52a960f09116c1e8ff50d4"


@pytest.fixture(scope="module")
def master():
    return load_master(MASTER)


def test_the_supplied_pilot_master_is_deployed(master):
    assert master["sha256"] == MASTER_SHA256
    assert (master["version"], master["rubric_version"], master["status"]) == ("1.0-PILOT", "SRT_Scoring_Rubric_v1", "PILOT")
    assert (master["decision_cap"], master["critical_failure_cap"], master["overall_max"]) == (6, 2, 90)
    assert (master["afi_yes_target"], master["afi_yes_floor"], master["second_review_proximity"]) == ([18, 22], 12, 2.0)


def test_bank_matches_the_workbook(master):
    srts = master["srts"]
    assert len(srts) == 100 and len({s["srt_id"] for s in srts}) == 100
    assert Counter(s["primary_competency"] for s in srts) == {f"C{i}": 10 for i in range(1, 11)}
    assert Counter(s["decision_applicable"] for s in srts) == {"Yes": 85, "No": 15}
    assert Counter(s["afi_applicability"] for s in srts) == {"Yes": 70, "Conditional": 28, "No": 2}
    assert Counter(s["critical_failure_severity"] for s in srts) == {"Major": 95, "Severe": 5}
    assert sorted(s["srt_id"] for s in srts if s["critical_failure_severity"] == "Severe") == [
        "C1-01", "C1-08", "C1-10", "C4-03", "C4-10"]
    assert "Material" not in {s["demand_type"] for s in srts}
    assert master["has_failure_type"] is False
    assert all(s["critical_failure_type"] is None for s in srts)


def test_competencies_bands_and_weights(master):
    comps = {c["code"]: c for c in master["competencies"]}
    assert [c["code"] for c in master["competencies"]] == [f"C{i}" for i in range(1, 11)]
    assert {k: (c["lens"], c["weight"]) for k, c in comps.items()} == {
        "C1": ("Business", 1.0), "C2": ("Business", 1.5), "C3": ("Business", 1.5), "C4": ("Business", 1.0),
        "C5": ("Business", 1.5), "C6": ("Technical", 1.0), "C7": ("Technical", 1.5), "C8": ("Technical", 1.0),
        "C9": ("Technical", 1.5), "C10": ("Technical", 1.5)}
    assert [(b["lower"], b["upper"], b["label"]) for b in master["readiness_bands"]] == [
        (0, 49, "Not Yet Ready"), (50, 62, "Developing"), (63, 75, "Ready"), (76, 90, "High Readiness")]
    assert [(b["lower"], b["upper"]) for b in master["afi_bands"]] == [(0, 39), (40, 59), (60, 74), (75, 89), (90, 100)]


def test_known_master_issues_are_reported_as_warnings(master):
    text = " ".join(master["warnings"])
    assert "Implementation_Decisions_2026-09" in text and "31" in text
    assert "Critical_Failure_Type" in text


# ── Refusals ─────────────────────────────────────────────────────────────────
def _variant(tmp_path, change):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(MASTER)
        change(wb)
        path = tmp_path / "variant.xlsx"
        wb.save(path)
    return path


def _column(ws, name):
    header = [c.value for c in ws[1]]
    return header.index(name) + 1


def _set(ws, name, srt_id, value):
    col, id_col = _column(ws, name), _column(ws, "SRT_ID")
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, id_col).value == srt_id:
            ws.cell(row, col).value = value
            return
    raise AssertionError(srt_id)


def _metadata(wb, field, value):
    ws = wb["Master_Metadata"]
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == field:
            ws.cell(row, 2).value = value
            return
    raise AssertionError(field)


@pytest.mark.parametrize("change,message", [
    (lambda wb: wb["PQI_SRT_Master_100"].delete_cols(_column(wb["PQI_SRT_Master_100"], "Decision_Applicable")),
     "lacks column(s): Decision_Applicable"),
    (lambda wb: _set(wb["PQI_SRT_Master_100"], "SRT_ID", "C1-02", "C1-01"), "duplicate SRT_ID(s): C1-01"),
    (lambda wb: _set(wb["PQI_SRT_Master_100"], "Anti_Firefighting_Applicability", "C2-01", "Maybe"),
     "Anti_Firefighting_Applicability 'Maybe'"),
    (lambda wb: _set(wb["PQI_SRT_Master_100"], "Situation", "C3-04", None), "Situation is empty"),
    (lambda wb: _metadata(wb, "Decision_Cap", None), "Decision_Cap is missing"),
    (lambda wb: _metadata(wb, "Assessment_Type", "PM"), "Assessment_Type must be PQI"),
    (lambda wb: _metadata(wb, "AFI_Yes_Emergency_Floor", "29"), "AFI=Yes floor of 29"),
    (lambda wb: [c for c in wb["Candidate_Output_Spec"]["A"] if c.value == "63–75"][0].__setattr__("value", "64–75"),
     "readiness bands are not contiguous"),
    (lambda wb: wb.remove(wb["AI_Evaluation_Sequence"]), "'AI_Evaluation_Sequence'"),
])
def test_unusable_master_is_refused(tmp_path, change, message):
    with pytest.raises(PqiMasterError) as info:
        load_master(_variant(tmp_path, change))
    assert message in str(info.value)


def test_missing_file_is_refused(tmp_path):
    with pytest.raises(PqiMasterError, match="not found"):
        load_master(tmp_path / "nope.xlsx")


def test_failure_type_column_is_used_when_present(tmp_path):
    def add_type(wb):
        ws = wb["PQI_SRT_Master_100"]
        col = ws.max_column + 1
        ws.cell(1, col).value = "Critical_Failure_Type"
        for row in range(2, ws.max_row + 1):
            ws.cell(row, col).value = "Technical-Safety" if ws.cell(row, 1).value in ("C4-03", "C4-10") else "Integrity"
    loaded = load_master(_variant(tmp_path, add_type))
    types = {s["srt_id"]: s["critical_failure_type"] for s in loaded["srts"]}
    assert loaded["has_failure_type"] is True
    assert types["C4-03"] == types["C4-10"] == "Technical-Safety" and types["C1-01"] == "Integrity"

    def bad_type(wb):
        add_type(wb)
        ws = wb["PQI_SRT_Master_100"]
        ws.cell(2, ws.max_column).value = "Safety"
    with pytest.raises(PqiMasterError, match="Critical_Failure_Type 'Safety'"):
        load_master(_variant(tmp_path, bad_type))
