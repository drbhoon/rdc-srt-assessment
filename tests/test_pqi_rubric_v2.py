"""SRT_Scoring_Rubric_v2: HR's revised PQI scoring engine (October 2026).

What must hold: one 0-9 level with the anti-bias guard on 0-3; the decision cap;
a Critical Failure capped only when the candidate's quoted words are really in
the answer; readiness and the second review on scale C1 with the pilot bands;
low confidence reviewed; completion figures kept apart from the official
score; every attempt traceable to its rubric, prompt, model and config; a full
re-score moving an attempt to v2 while keeping the result it replaced; and the
validation workbook. Rubric v1 itself is covered, unchanged, by the other
PQI test files.
"""
import copy
import io
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import main
import pqi_evaluator as ev
import pqi_fakes as pf
import pqi_rubric_v2 as v2
import pqi_scoring as sc
import pqi_service
from test_pm_api import ADMIN
from test_pqi_flow import answer_all, form_with, new_code, session, start, use_claude
from test_pqi_master import master  # noqa: F401


def srt(master, srt_id, **changes):
    record = copy.deepcopy(next(s for s in master["srts"] if s["srt_id"] == srt_id))
    record.update(changes)
    return record


def reply(**changes):
    base = {"srt_id": "C1-01", "raw_level": 6, "decision_made": "Yes", "essential_judgment": "Yes",
            "key_evidence": ["Refused the inducement and told the BH."], "material_gap": "No recurrence control.",
            "low_score_basis": v2.NOT_LOW, "serious_negative": "No", "critical_failure": "No",
            "critical_failure_statement": "", "critical_failure_reason": "", "unresolved_contradiction": False,
            "afi_activated": "Yes", "afi_evidence_level": 2, "afi_evidence": "Fixed today's case only.",
            "evaluation_confidence": "High"}
    base.update(changes)
    return base


def observed(**changes):
    return {**ev.parse_evaluation_v2(SRT_YES, reply()), **changes}


SRT_YES = {"srt_id": "C1-01", "decision_applicable": "Yes", "afi_applicability": "Yes"}


# ── The evaluator's record ───────────────────────────────────────────────────
def test_a_valid_v2_evaluation_is_normalised(master):
    parsed = ev.parse_evaluation_v2(srt(master, "C1-01"), reply(raw_level="6", decision_made="yes",
                                                              serious_negative="no", evaluation_confidence="low"))
    assert (parsed["raw_level"], parsed["decision_made"], parsed["serious_negative"],
            parsed["evaluation_confidence"]) == (6, "Yes", False, "Low")


@pytest.mark.parametrize("changes,message", [
    ({"srt_id": "C1-02"}, "does not match"),
    ({"raw_level": 10}, "raw_level must be 0–9"),
    ({"raw_level": 3}, "must give the reason for a level of 0–3"),
    ({"raw_level": 5, "low_score_basis": "Missed the essential issue"}, "for a level of 4 or higher"),
    ({"low_score_basis": "Answer was short"}, "low_score_basis must be one of"),
    ({"decision_made": "Not Applicable"}, "requires a decision"),
    ({"critical_failure": "Yes", "critical_failure_statement": " "}, "must quote the candidate's words"),
    ({"afi_evidence_level": None}, "afi_evidence_level is required"),
    ({"key_evidence": []}, "key_evidence"),
    ({"essential_judgment": "Mostly"}, "essential_judgment must be one of"),
    ({"evaluation_confidence": "Unsure"}, "evaluation_confidence must be one of"),
])
def test_malformed_v2_evaluations_are_rejected(master, changes, message):
    with pytest.raises(ev.ReplyFormatError, match=message):
        ev.parse_evaluation_v2(srt(master, "C1-01", decision_applicable="Yes", afi_applicability="Yes"),
                               reply(**changes))


def test_no_level_of_0_to_3_without_one_of_the_four_reasons():
    assert set(v2.LOW_SCORE_BASES) == {"Missed the essential issue", "Materially deficient judgment",
                                       "Advocated an unsafe or wrong action", "No usable answer"}
    tool = ev.EVALUATION_TOOL_V2["input_schema"]
    assert set(tool["required"]) == set(tool["properties"])
    # Severity and type stay the master's: the evaluator is not asked for them.
    assert not any("severity" in k or "type" in k for k in tool["properties"])


# ── What the evaluator is given ──────────────────────────────────────────────
def test_v2_prompt_replaces_the_v1_rulebook(master):
    v2_prompt = ev.evaluator_system_prompt(master, ev.RUBRIC_V2)
    v1_prompt = ev.evaluator_system_prompt(master)
    assert v2.FINAL_PRINCIPLE in v2_prompt
    assert "Level 5 means a competent PQI could reasonably act on this answer." in v2_prompt
    assert "## Sheet: Competency_Framework" in v2_prompt and "## Sheet: AFI_Rubric" in v2_prompt
    for sheet in ("SRT_Scoring_Rubric_v1", "AI_Scoring_Framework", "AI_Evaluation_Sequence",
                  "AI_Evaluator_Output_Schema", "Scoring_Architecture", "Master_Metadata"):
        assert f"## Sheet: {sheet}" not in v2_prompt and f"## Sheet: {sheet}" in v1_prompt
    assert v2.RUBRIC_TEXT not in v1_prompt
    # Nothing in the rubric is about a particular person or group.
    for word in ("Haseeb", "designation bonus", "experience bonus"):
        assert word not in v2_prompt

    block = ev._srt_block(master, srt(master, "C7-03"), ev.RUBRIC_V2)
    assert "Critical_Failure_Control" not in block and "Scoring_Note" not in block
    assert block["Model_Response"] == srt(master, "C7-03")["model_response"]
    assert ev.prompt_version(master, ev.RUBRIC_V2).startswith(v2.PROMPT_TEMPLATE + ":")
    assert ev.prompt_version(master) == ev.prompt_version(master, ev.RUBRIC_V1)
    assert ev.prompt_version(master).startswith("pqi-evaluator-2:")
    assert v2.REPORT_GUIDANCE in ev.report_system_prompt(master, ev.RUBRIC_V2)
    assert v2.REPORT_GUIDANCE not in ev.report_system_prompt(master)


# ── Rules applied in code ────────────────────────────────────────────────────
def test_the_level_stands_unless_a_cap_applies(master):
    record = srt(master, "C1-01", decision_applicable="Yes")
    assert sc.apply_rules_v2(master, record, observed(raw_level=9))["final_score"] == 9
    capped = sc.apply_rules_v2(master, record, observed(raw_level=8, decision_made="No"))
    assert (capped["final_score"], capped["caps_applied"], capped["raw_level"]) == (6, ["decision"], 8)
    assert capped["calibrated_percentage"] == 80.0
    assert sc.apply_rules_v2(master, record, observed(raw_level=5, decision_made="No"))["final_score"] == 5


@pytest.mark.parametrize("quote,answer,verified", [
    ("I will accept the gift", "Honestly, I will accept the gift and keep it quiet.", True),
    ("i will ACCEPT the gift!", "Honestly — I will accept the gift, and keep it quiet.", True),
    ("we add water at site to get slump back", "We add some water at site to get the slump back.", True),
    ("I would report it to the BH immediately", "I will refuse and inform QA.", False),
    ("accept", "I will not accept anything.", True),
    ("", "anything", False),
])
def test_a_critical_failure_must_quote_the_candidate(quote, answer, verified):
    assert sc.statement_in_response(quote, answer) is verified


def test_critical_failure_is_capped_only_when_the_quote_is_real(master):
    record = srt(master, "C1-01")
    answer = "Take the payment from the supplier, nobody will know."
    confirmed = sc.apply_rules_v2(master, record, observed(
        raw_level=5, critical_failure=True, critical_failure_statement="take the payment from the supplier",
        critical_failure_reason="Accepts an inducement."), answer)
    assert (confirmed["final_score"], confirmed["critical_failure"], confirmed["caps_applied"]) == (
        2, True, ["critical_failure"])
    assert confirmed["critical_failure_severity"] == record["critical_failure_severity"]

    paraphrased = sc.apply_rules_v2(master, record, observed(
        raw_level=5, critical_failure=True, critical_failure_statement="The candidate would accept a bribe",
        critical_failure_reason="Accepts an inducement."), "I will refuse it and tell my BH.")
    assert (paraphrased["final_score"], paraphrased["critical_failure"]) == (5, False)
    assert paraphrased["critical_failure_unverified"] is True
    assert paraphrased["critical_failure_claim"]["statement"] == "The candidate would accept a bribe"


def test_blank_answers_under_v2(master):
    blank = sc.blank_result(master, srt(master, "C1-01", afi_applicability="Yes"), ev.RUBRIC_V2)
    assert (blank["status"], blank["final_score"], blank["afi_evidence_level"], blank["rubric"]) == (
        "blank", 0, 0, ev.RUBRIC_V2)
    conditional = sc.blank_result(master, srt(master, "C1-01", afi_applicability="Conditional"), ev.RUBRIC_V2)
    assert (conditional["afi_activated"], conditional["afi_evidence_level"]) == (False, None)


def _results(master, levels):
    """Three SRTs per competency at the given levels (one level, or 30)."""
    levels = levels if isinstance(levels, list) else [levels] * 30
    out = []
    for i, level in enumerate(levels):
        code = f"C{i // 3 + 1}"
        out.append({"srt_id": f"{code}-{i % 3 + 1:02d}", "primary_competency": code, "status": "scored",
                    "final_score": level, "critical_failure": False, "unresolved_contradiction": False,
                    "afi_activated": False, "afi_evidence_level": None, "critical_failure_severity": "None",
                    "critical_failure_type": "None"})
    return out


@pytest.mark.parametrize("levels,band", [
    (6, "High Readiness"),            # every SRT 80% → 80.0
    (5, "Ready"),                     # 70.0
    (4, "Developing"),                # 60.0
    (3, "Developing"),                # 50.0
    (2, "Not Yet Ready"),             # 35.0
])
def test_v2_readiness_uses_the_pilot_bands_on_c1(master, levels, band):
    results = _results(master, levels)
    decided = sc.basis(master, ev.RUBRIC_V2, results)
    ready = sc.readiness(master, sc.headline(master, results), results, [], decided)
    assert ready["band"] == band
    assert ready["decided_on"] == round(sc.reported_headline(master, results)["overall"], 1)
    assert [b["lower"] for b in ready["bands"]] == [0, 50, 65, 80]
    assert ready["provisional_note"] == v2.PILOT_NOTE


def test_v2_second_review_is_triggered_on_c1_and_by_low_confidence(master):
    # 30 x level 5 = 70.0 on C1: not within 2 of 65 or 80.
    results = _results(master, 5)
    decided = sc.basis(master, ev.RUBRIC_V2, results)
    assert sc.review_plan(master, sc.headline(master, results), results, [], decided)["triggers"] == []
    # Swap three 5s for 4s: 69.0 ... and two more → 67.0 → within 2 of 65.
    near = _results(master, [4] * 9 + [5] * 21)
    plan = sc.review_plan(master, sc.headline(master, near), near, [], sc.basis(master, ev.RUBRIC_V2, near))
    assert "near_readiness_boundary" in plan["triggers"] and len(plan["srts"]) == 30

    results[4]["evaluation_confidence"] = "Low"
    results[7]["critical_failure_unverified"] = True
    plan = sc.review_plan(master, sc.headline(master, results), results, [], decided)
    assert set(plan["triggers"]) == {"low_confidence", "unverified_critical_failure"}
    assert set(plan["srts"]) == {results[4]["srt_id"], results[7]["srt_id"]}
    for reason in plan["srts"].values():
        assert "overall" not in reason.lower() and "band" not in reason.lower()


def test_completion_is_kept_apart_from_the_score(master):
    results = _results(master, 6)
    for i in (0, 1, 2, 10):
        results[i].update(status="blank", final_score=0)
    done = sc.completion(master, results)
    assert (done["questions_presented"], done["questions_answered"], done["blank_count"]) == (30, 26, 4)
    assert done["completion_percent"] == 86.7
    assert done["answered_only_overall"] == 80.0          # C1 has no answer, so it is left out
    assert sc.reported_headline(master, results)["overall"] < 80.0


# ── End to end ───────────────────────────────────────────────────────────────
USAGE = SimpleNamespace(input_tokens=1200, output_tokens=2500, cache_creation_input_tokens=0,
                        cache_read_input_tokens=9000)


@pytest.fixture
def pqi2(store, monkeypatch):
    monkeypatch.setattr(pqi_service, "CURRENT_RUBRIC", ev.RUBRIC_V2)
    fake = pf.FakePqiClaude(usage=USAGE)
    monkeypatch.setattr(main, "client", fake)
    monkeypatch.setattr(ev, "time", SimpleNamespace(sleep=lambda s: None))
    with TestClient(main.app) as client:
        client.fake = fake
        client.store = store
        yield client


def test_a_v2_attempt_is_scored_and_recorded_under_v2(pqi2):
    sid = start(pqi2, new_code(pqi2)).json()["session_id"]
    stored = session(sid)
    assert stored["rubric_version"] == ev.RUBRIC_V2
    assert stored["prompt_version"].startswith(v2.PROMPT_TEMPLATE + ":")

    answer_all(pqi2, sid, blank=(3,))
    stored = session(sid)
    report, master = stored["report"], pqi_service.current_master()
    assert stored["status"] == "completed"
    assert {c["tool"] for c in pqi2.fake.calls if c["mode"] == "evaluate"} == {"record_srt_evaluation_v2"}

    ordered = [stored["scores"][s["srt_id"]] for s in stored["srt_snapshot"]]
    assert all(r["rubric"] == ev.RUBRIC_V2 for r in ordered)
    assert (report["scoring_rubric"], report["evaluation_mode"], report["calibration"]) == (
        ev.RUBRIC_V2, "PILOT_CALIBRATION", "C1")
    assert len(report["config_hash"]) == 12 and report["question_ids"] == [s["srt_id"] for s in stored["srt_snapshot"]]
    assert report["prompt_version"] == stored["prompt_version"]
    assert (report["completion"]["blank_count"], report["completion"]["questions_answered"]) == (1, 29)

    decided = sc.basis(master, ev.RUBRIC_V2, ordered)
    assert report["overall_pqi_score"] == round(decided["overall"], 1)
    assert report["readiness"]["basis"] == v2.PILOT_BANDS_LABEL
    assert report["overall_readiness"] == sc.readiness(master, sc.headline(master, ordered), ordered, [],
                                                       decided)["band"]
    assert report["headline"]["rubric"] == ev.RUBRIC_V2 and report["headline"]["blank_count"] == 1

    usage = report["usage"]
    assert usage["calls"] == len(pqi2.fake.calls)
    assert usage["output_tokens"] == 2500 * len(pqi2.fake.calls) and usage["estimated_cost_usd"] > 0

    # The evaluator saw the SRT and the answer — never who answered or how long they took.
    call = next(c for c in pqi2.fake.calls if c["mode"] == "evaluate")
    told = str(call["payload"]).lower()
    for secret in ("outside applicant", "designation", "minutes", "deadline", "word count"):
        assert secret not in told
    assert "Critical_Failure_Control" not in call["payload"]["SRT"]
    narrative = next(c for c in pqi2.fake.calls if c["mode"] == "report")["payload"]
    assert narrative["Scoring_Rubric"] == ev.RUBRIC_V2 and narrative["Completion"]["blank_count"] == 1

    pdf = pqi2.get(f"/api/admin/download-pdf/{sid}", headers=ADMIN)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_an_unsupported_critical_failure_is_reviewed_not_capped(pqi2, monkeypatch):
    form_with(monkeypatch, "C1-01", "C1-02")

    def rule(srt_id, payload):
        if srt_id == "C1-01":       # quotes words the candidate never wrote
            return {"level": 5, "critical": "Yes", "statement": "I would take the money",
                    "reason": "Accepts an inducement."}
        if srt_id == "C1-02":       # quotes the candidate's own words
            return {"level": 5, "critical": "Yes", "statement": "hold the load, test, record",
                    "reason": "Stand-in for a confirmed failure."}
        return {"level": 5, "decision": "Yes"}
    use_claude(pqi2, monkeypatch, pf.FakePqiClaude(evaluate_v2=pf.evaluator_v2(rule)))
    sid = start(pqi2, new_code(pqi2)).json()["session_id"]
    answer_all(pqi2, sid)
    stored = session(sid)
    unsupported, confirmed = stored["scores"]["C1-01"], stored["scores"]["C1-02"]

    assert (unsupported["final_score"], unsupported["critical_failure"], unsupported["critical_failure_unverified"]) == (
        5, False, True)
    assert (confirmed["final_score"], confirmed["critical_failure"]) == (2, True)
    assert {"C1-01", "C1-02"} <= set(pqi2.fake.srt_ids("review"))
    report = stored["report"]
    assert "unverified_critical_failure" in report["second_review"]["triggers"]
    assert [f["srt_id"] for f in report["critical_flags"]] == ["C1-02"]
    assert report["critical_flags"][0]["statement"] == "hold the load, test, record"
    assert any("C1-01" in reason for reason in report["readiness"]["manual_review_reasons"])
    # The reviewer is shown what was claimed, and why it is being asked again.
    review = next(c for c in pqi2.fake.calls if c["mode"] == "review" and c["payload"]["SRT"]["SRT_ID"] == "C1-01")
    assert review["payload"]["Earlier_Evaluation"]["Critical_Failure_Statement"] == "I would take the money"
    assert "could not be found" in review["instruction"]


def test_low_confidence_triggers_a_second_review(pqi2, monkeypatch):
    form_with(monkeypatch, "C2-01")

    def rule(srt_id, payload):
        return {"level": 5, "confidence": "Low" if srt_id == "C2-01" else "High", "decision": "Yes"}
    use_claude(pqi2, monkeypatch, pf.FakePqiClaude(evaluate_v2=pf.evaluator_v2(rule)))
    sid = start(pqi2, new_code(pqi2)).json()["session_id"]
    answer_all(pqi2, sid)
    report = session(sid)["report"]
    assert "low_confidence" in report["second_review"]["triggers"]
    assert "C2-01" in pqi2.fake.srt_ids("review")


def test_full_rescore_moves_a_v1_attempt_to_v2_and_keeps_its_old_result(pqi2, monkeypatch):
    monkeypatch.setattr(pqi_service, "CURRENT_RUBRIC", ev.RUBRIC_V1)
    v1_sid = start(pqi2, new_code(pqi2)).json()["session_id"]
    answer_all(pqi2, v1_sid)
    before = copy.deepcopy(session(v1_sid))       # the memory store hands back the live record
    assert before["rubric_version"] == ev.RUBRIC_V1 and before["report"].get("scoring_rubric") == ev.RUBRIC_V1

    monkeypatch.setattr(pqi_service, "CURRENT_RUBRIC", ev.RUBRIC_V2)
    v2_sid = start(pqi2, new_code(pqi2)).json()["session_id"]
    answer_all(pqi2, v2_sid)
    use_claude(pqi2, monkeypatch, pf.FakePqiClaude(usage=USAGE))

    queued = pqi2.post("/api/admin/pqi/rescore-all", headers=ADMIN).json()
    assert [q["session_id"] for q in queued["queued"]] == [v1_sid]
    assert any(s["session_id"] == v2_sid and "already on" in s["reason"] for s in queued["skipped"])
    assert {c["tool"] for c in pqi2.fake.calls if c["mode"] == "evaluate"} == {"record_srt_evaluation_v2"}

    after = session(v1_sid)
    assert (after["status"], after["rubric_version"]) == ("completed", ev.RUBRIC_V2)
    assert after["report"]["scoring_rubric"] == ev.RUBRIC_V2
    assert after["prompt_version"].startswith(v2.PROMPT_TEMPLATE + ":")
    [old] = after["report_history"]
    assert (old["scoring_rubric"], old["overall"], old["readiness"]) == (
        ev.RUBRIC_V1, before["report"]["overall_pqi_score"], before["report"]["overall_readiness"])
    assert old["srt_levels"] == {r["srt_id"]: r["final_score"] for r in before["report"]["srt_results"]}

    evaluations = pqi2.get(f"/api/admin/session-evaluations/{v1_sid}", headers=ADMIN).json()["evaluations"]
    assert {e["rubric_version"] for e in evaluations} == {ev.RUBRIC_V1, ev.RUBRIC_V2}

    # The single-attempt Full button does the same.
    full = pqi2.post(f"/api/admin/rescore/{v2_sid}?force_full=true", headers=ADMIN)
    assert full.status_code == 200
    assert len(session(v2_sid)["report_history"]) == 1

    # The validation table.
    xlsx = pqi2.get("/api/admin/pqi/validation.xlsx", headers=ADMIN)
    assert xlsx.status_code == 200 and "PQI_Validation_" in xlsx.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Validation", "Competencies", "SRT levels", "Levels by question",
                             "Level distribution", "Critical failures", "Second review", "Read me"]
    header = [c.value for c in wb["Validation"][1]]
    for column in ("Candidate", "Previous Score", "Revised Score", "Technical", "Business", "AFI", "Readiness",
                   "Blank Count"):
        assert column in header
    rows = [dict(zip(header, [c.value for c in row])) for row in wb["Validation"].iter_rows(min_row=2)]
    moved = next(r for r in rows if r["Previous Rubric"] == ev.RUBRIC_V1)
    assert moved["Previous Score"] == before["report"]["overall_pqi_score"]
    assert moved["Revised Score"] == after["report"]["overall_pqi_score"]
    assert moved["Rubric"] == ev.RUBRIC_V2
    distribution = list(wb["Level distribution"].iter_rows(min_row=2, values_only=True))
    assert all(sum(row[1:]) == 30 for row in distribution)
    assert wb["SRT levels"].max_row == 1 + 30 * len(rows)

    assert pqi2.get("/api/admin/pqi/validation.xlsx").status_code == 401
    assert pqi2.post("/api/admin/pqi/rescore-all").status_code == 401


def test_reissue_keeps_a_v2_report_on_v2(pqi2, monkeypatch):
    sid = start(pqi2, new_code(pqi2)).json()["session_id"]
    answer_all(pqi2, sid)
    before = copy.deepcopy(session(sid))
    use_claude(pqi2, monkeypatch, pf.FakePqiClaude())
    assert pqi2.post(f"/api/admin/pqi/reissue/{sid}", headers=ADMIN).status_code == 200
    after = session(sid)
    assert pqi2.fake.modes() == ["report"]
    assert after["report"]["scoring_rubric"] == ev.RUBRIC_V2
    assert after["report"]["overall_readiness"] == before["report"]["overall_readiness"]
    assert after["report"]["readiness"]["basis"] == v2.PILOT_BANDS_LABEL


def test_a_rejected_reply_is_retried_with_the_reason(pqi2, monkeypatch):
    # Production, 2026-10-08: the evaluator left AFI_Evidence_Level empty on
    # AFI=Yes SRTs, the identical retry got the identical reply, and the run
    # failed three times over. The retry now says what was wrong.
    form_with(monkeypatch, "C1-10")
    seen = {"n": 0}

    def rule(srt_id, payload):
        if srt_id == "C1-10":
            seen["n"] += 1
            if seen["n"] == 1:
                return {"level": 4, "afi_level": None}
        return {"level": 5, "decision": "Yes"}
    use_claude(pqi2, monkeypatch, pf.FakePqiClaude(evaluate_v2=pf.evaluator_v2(rule)))
    sid = start(pqi2, new_code(pqi2)).json()["session_id"]
    answer_all(pqi2, sid)
    stored = session(sid)

    assert stored["status"] == "completed"
    calls = [c for c in pqi2.fake.calls if c["mode"] == "evaluate" and c["payload"]["SRT"]["SRT_ID"] == "C1-10"]
    assert len(calls) == 2
    assert not calls[0]["instruction"].startswith("Your previous")
    told = calls[1]["instruction"]
    assert told.startswith("Your previous record_srt_evaluation_v2 call was rejected: afi_evidence_level is required")
    assert "Anti_Firefighting_Applicability is Yes" in told and "using 0 when" in told
    # Same system prompt (cached) and same prompt version as every other call.
    assert calls[1]["kwargs"]["system"] == calls[0]["kwargs"]["system"]
    assert stored["scores"]["C1-10"]["afi_evidence_level"] is not None
    assert stored["report"]["prompt_version"] == stored["prompt_version"]
