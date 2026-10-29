"""The Plant Manager results report for HR — same layout as PQI's, from what the pipeline stored."""
import copy
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import database
import main
import pm_fakes as fk
from question_bank import REQUIRED_COMPETENCIES
from test_pm_api import ADMIN
from test_pm_pipeline import SESSION_ID, run, sleeps  # noqa: F401  (fixture)
from test_pqi_results import rows, set_employee_code


@pytest.fixture
def api(store, monkeypatch):
    looked_up = []

    def designations(codes):
        looked_up.append(sorted(c for c in codes if c))
        return {"g00064": "Plant Incharge"}
    monkeypatch.setattr(main, "master_designations", designations)
    with TestClient(main.app) as client:
        client.store, client.looked_up = store, looked_up
        yield client


def test_plant_manager_results_report(api, monkeypatch, sleeps):
    qs = fk.questions(seed=101)
    stored = run(monkeypatch, fk.FakeClaude(score=fk.scorer(5, 8), report=fk.reporter()), qs,
                 fk.answers(qs))
    assert stored["status"] == "completed"
    report = copy.deepcopy(stored["report"])
    set_employee_code(api, SESSION_ID, "G00064")
    database.create_session("pm-waiting", dict(fk.CANDIDATE, candidate_name="Still Answering"), qs)
    main._cache.clear()

    r = api.get("/api/admin/pm/results.xlsx", headers=ADMIN)
    assert r.status_code == 200 and 'filename="PlantManager_Results_' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Results", "Not completed", "Competency key", "How to read"]

    header, results = rows(wb["Results"], 4)
    for column in ("Emp Code", "Name", "Designation", "Location", "Total /300", "Score /100", "Readiness",
                   "Integrity flag", "Answered", *REQUIRED_COMPETENCIES, "Top strengths", "Development areas"):
        assert column in header
    [row] = results
    assert (row["Emp Code"], row["Designation"]) == ("G00064", "Plant Incharge")
    assert row["Name"] == fk.CANDIDATE["candidate_name"]
    assert (row["Total /300"], row["Score /100"], row["Base total /300 (before English)"]) == (
        report["overall_score_out_of_300"], report["normalized_score_out_of_100"], report["base_score_out_of_300"])
    assert row["Readiness"] == report["overall_readiness"]
    for name in REQUIRED_COMPETENCIES:
        assert row[name] == report["competency_summary"][name]
    if report.get("readiness_demoted_from"):
        assert row["Readiness note"].startswith(f"Lowered from {report['readiness_demoted_from']}")
    assert row["Top strengths"].split("; ")[0] == report["top_strengths"][0]["strength"]
    assert api.looked_up == [["G00064"]]

    _, pending = rows(wb["Not completed"], 1)
    assert [p["Name"] for p in pending] == ["Still Answering"]
    tiers = "\n".join(c.value or "" for c in wb["How to read"]["A"])
    assert "Ready to be Plant Manager: total 210+ /300 and every competency at least 6/10" in tiers
    assert wb["Competency key"].max_row == 11

    # Plant Manager attempts never appear in the PQI report, and the reverse.
    _, pqi_rows = rows(load_workbook(io.BytesIO(api.get("/api/admin/pqi/results.xlsx", headers=ADMIN).content))
                       ["Results"], 4)
    assert pqi_rows == []
    assert api.get("/api/admin/pm/results.xlsx").status_code == 401
