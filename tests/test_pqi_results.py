"""The PQI results report for HR: one row per candidate, from what the database holds."""
import copy
import io
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import database
import main
import pqi_evaluator as ev
import pqi_fakes as pf
import pqi_service
from test_pm_api import ADMIN
from test_pqi_flow import answer_all, new_code, session, start


@pytest.fixture
def api(store, monkeypatch):
    monkeypatch.setattr(pqi_service, "CURRENT_RUBRIC", ev.RUBRIC_V2)
    monkeypatch.setattr(main, "client", pf.FakePqiClaude())
    monkeypatch.setattr(ev, "time", SimpleNamespace(sleep=lambda s: None))
    looked_up = []

    def designations(codes):
        looked_up.append(sorted(c for c in codes if c))
        return {"a00252": "Senior Manager - Quality"}
    monkeypatch.setattr(main, "master_designations", designations)
    with TestClient(main.app) as client:
        client.store, client.looked_up = store, looked_up
        yield client


def set_employee_code(api, sid, code):
    if api.store == "postgres":
        conn = database._get_conn()
        try:
            with conn.cursor() as cur:
                cur.execute("UPDATE sessions SET employee_code = %s WHERE session_id = %s", (code, sid))
            conn.commit()
        finally:
            conn.close()
    else:
        database._memory_store[sid]["candidate"]["employee_code"] = code
    main._cache.clear()


def rows(sheet, header_row):
    header = [c.value for c in sheet[header_row]]
    return header, [dict(zip(header, [c.value for c in r])) for r in sheet.iter_rows(min_row=header_row + 1)]


def test_results_report_has_one_row_per_candidate(api):
    done = start(api, new_code(api)).json()["session_id"]
    answer_all(api, done, blank=(5,))
    set_employee_code(api, done, "A00252")
    waiting = start(api, new_code(api), body={
        "candidate_name": "Still Answering", "plant_location": "Pune", "assessment_date": "2026-10-09",
        "candidate_type": "external", "email": "second@example.com"}).json()["session_id"]
    report = copy.deepcopy(session(done)["report"])

    r = api.get("/api/admin/pqi/results.xlsx", headers=ADMIN)
    assert r.status_code == 200 and 'filename="PQI_Results_' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Results", "Not completed", "Competency key", "How to read"]

    header, results = rows(wb["Results"], 4)
    for column in ("Emp Code", "Name", "Designation", "Location", "Overall /100", "Technical /100", "Business /100",
                   "AFI /100", "Readiness", "Answered", "C1 Integrity & Trust",
                   "C10 Functional Knowledge, Data Analysis & Technical Judgment"):
        assert column in header
    [row] = results
    assert (row["Emp Code"], row["Designation"], row["Name"]) == ("A00252", "Senior Manager - Quality",
                                                                  "Outside Applicant")
    assert row["Overall /100"] == report["overall_pqi_score"]
    assert (row["Technical /100"], row["Business /100"], row["AFI /100"]) == (
        report["technical_acumen"], report["business_acumen"], report["anti_firefighting_index"])
    assert row["Readiness"] == report["overall_readiness"] and row["Answered"] == "29/30"
    scores = {c["code"]: c["score"] for c in report["competency_scores"]}
    assert row["C1 Integrity & Trust"] == scores["C1"]
    assert row["Scoring rubric"] == "v2" and row["E-mail"] == "applicant@example.com"
    assert api.looked_up == [["A00252"]]

    _, pending = rows(wb["Not completed"], 1)
    assert [(p["Name"], p["Status"]) for p in pending] == [("Still Answering", "Taking the test")]
    assert waiting

    assert wb["Competency key"].max_row == 11
    assert api.get("/api/admin/pqi/results.xlsx").status_code == 401


def test_results_report_without_the_master_or_any_candidate(api, monkeypatch):
    monkeypatch.setattr(main, "master_designations", lambda codes: {})
    r = api.get("/api/admin/pqi/results.xlsx", headers=ADMIN)
    wb = load_workbook(io.BytesIO(r.content))
    header, results = rows(wb["Results"], 4)
    assert results == [] and "Readiness" in header
