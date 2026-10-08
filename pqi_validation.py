"""The PQI validation workbook: every candidate's result under the current
scoring engine, beside the result it replaced.

HR's engine review (October 2026) asked for this after any rubric change: the
headline table (previous score, revised score, technical, business, AFI,
readiness, blank count) and the detail behind it — competency scores, all 30
SRT levels, the count of each level 0-9, Critical Failures and second-review
changes — so the revised engine can be checked against expert judgment
without rebuilding anything from PDFs.

Read-only: built from what the database already holds, no AI calls.
"""
from __future__ import annotations

import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
COMPETENCIES = [f"C{i}" for i in range(1, 11)]


def _previous(session: dict, report: dict) -> dict:
    """The result this report replaced: the latest earlier report on a
    different rubric or scale, else simply the latest earlier report."""
    history = session.get("report_history") or []
    current = (report.get("scoring_rubric"), (report.get("score_scale") or {}).get("name"))
    for entry in reversed(history):
        if (entry.get("scoring_rubric"), entry.get("score_scale")) != current:
            return entry
    return history[-1] if history else {}


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "B2"
    for i, h in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(40, len(str(h)) + 2)))
    return ws


def build_workbook(sessions: list) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    summary, competencies, levels_long, levels_wide, distribution, failures, reviews = [], [], [], [], [], [], []
    ordered = sorted(sessions, key=lambda s: -(((s.get("report") or {}).get("overall_pqi_score")) or -1))
    for s in ordered:
        name = (s.get("candidate") or {}).get("candidate_name") or s.get("session_id")
        report = s.get("report") or {}
        if not report:
            summary.append([name, (s.get("candidate") or {}).get("plant_location"), s.get("status")] + [None] * 22)
            continue
        prev = _previous(s, report)
        done = report.get("completion") or {}
        review = report.get("second_review") or {}
        results = report.get("srt_results") or []
        usage = report.get("usage") or {}
        summary.append([
            name, (s.get("candidate") or {}).get("plant_location"), s.get("status"),
            prev.get("overall"), prev.get("scoring_rubric"), prev.get("score_scale"), prev.get("readiness"),
            report.get("overall_pqi_score"), report.get("technical_acumen"), report.get("business_acumen"),
            report.get("anti_firefighting_index"), report.get("overall_readiness"),
            done.get("blank_count"), done.get("completion_percent"), done.get("answered_only_overall"),
            len(report.get("critical_flags") or []), len(review.get("changes") or []),
            ", ".join(review.get("triggers") or []),
            "Yes" if (report.get("readiness") or {}).get("manual_review_required") else "No",
            report.get("scoring_rubric") or "SRT_Scoring_Rubric_v1",
            (report.get("score_scale") or {}).get("name"), report.get("prompt_version"),
            report.get("evaluator_model"), report.get("config_hash"), usage.get("estimated_cost_usd"),
        ])
        by_code = {c["code"]: c.get("score") for c in report.get("competency_scores") or []}
        competencies.append([name] + [by_code.get(c) for c in COMPETENCIES])

        counts = Counter(r.get("final_score") for r in results)
        distribution.append([name] + [counts.get(level, 0) for level in range(10)])
        levels_wide.append([name] + [r.get("final_score") for r in results])
        for r in results:
            first = (r.get("first_pass") or {}).get("final_score")
            levels_long.append([
                name, r.get("question_number"), r.get("srt_id"), r.get("primary_competency"),
                r.get("raw_level", r.get("core_response_level")), r.get("final_score"),
                r.get("calibrated_percentage"), ", ".join(r.get("caps_applied") or []), r.get("decision_made"),
                r.get("essential_judgment"), r.get("evaluation_confidence"),
                "Yes" if r.get("serious_negative") else ("No" if "serious_negative" in r else None),
                "Yes" if r.get("critical_failure") else "No",
                r.get("afi_evidence_level") if r.get("afi_activated") else None,
                "Yes" if r.get("reviewed") else "No", first if r.get("reviewed") else None,
                r.get("status"), r.get("low_score_basis"), r.get("primary_gap"),
            ])
            if r.get("critical_failure"):
                failures.append([name, r.get("srt_id"), "Confirmed", r.get("critical_failure_severity"),
                                 r.get("critical_failure_type"), r.get("critical_failure_statement"),
                                 r.get("critical_failure_reason")])
            elif r.get("critical_failure_unverified"):
                claim = r.get("critical_failure_claim") or {}
                failures.append([name, r.get("srt_id"), "Not confirmed (quote not found in the answer)",
                                 None, None, claim.get("statement"), claim.get("reason")])
        for change in review.get("changes") or []:
            reviews.append([name, change.get("srt_id"), change.get("from_score"), change.get("to_score"),
                            "Yes" if change.get("from_critical_failure") else "No",
                            "Yes" if change.get("to_critical_failure") else "No"])

    _sheet(wb, "Validation", [
        "Candidate", "Location", "Status", "Previous Score", "Previous Rubric", "Previous Scale",
        "Previous Readiness", "Revised Score", "Technical", "Business", "AFI", "Readiness", "Blank Count",
        "Completion %", "Answered-only (diagnostic)", "Critical Failures", "Second-review changes",
        "Second-review triggers", "Manual review", "Rubric", "Calibration", "Prompt version", "Evaluator model",
        "Config hash", "Est. cost (USD)",
    ], summary, {"Candidate": 26, "Location": 18, "Second-review triggers": 40, "Prompt version": 30,
                 "Rubric": 24, "Previous Rubric": 24})
    _sheet(wb, "Competencies", ["Candidate"] + COMPETENCIES, competencies, {"Candidate": 26})
    _sheet(wb, "SRT levels", [
        "Candidate", "Q", "SRT_ID", "Competency", "Raw level", "Final level", "C1 %", "Caps", "Decision made",
        "Essential judgment", "Confidence", "Serious negative", "Critical failure", "AFI level", "Reviewed",
        "First-pass level", "Status", "Low-score basis", "Material gap",
    ], levels_long, {"Candidate": 26, "Material gap": 60, "Low-score basis": 30})
    _sheet(wb, "Levels by question", ["Candidate"] + [f"Q{i}" for i in range(1, 31)], levels_wide,
           {"Candidate": 26})
    _sheet(wb, "Level distribution", ["Candidate"] + [f"Level {i}" for i in range(10)], distribution,
           {"Candidate": 26})
    _sheet(wb, "Critical failures", ["Candidate", "SRT_ID", "Status", "Severity", "Type", "Candidate's words",
                                     "Reason"], failures, {"Candidate": 26, "Candidate's words": 60, "Reason": 60})
    _sheet(wb, "Second review", ["Candidate", "SRT_ID", "From level", "To level", "Critical failure before",
                                 "Critical failure after"], reviews, {"Candidate": 26})
    notes = wb.create_sheet("Read me")
    for line in [
        "PQI validation workbook — built from the database, no AI calls.",
        "Revised Score, Technical, Business and competency scores are percentages on the report's calibration "
        "(C1: level 0-9 = 0, 20, 35, 50, 60, 70, 80, 88, 94, 100), averaged within each competency.",
        "Previous Score is the latest earlier report on a different rubric or scale, as it was printed.",
        "Blank answers score 0 in the official result. Answered-only is diagnostic and is never the score.",
        "Compare candidates only when Rubric, Prompt version and Evaluator model match (Config hash equal).",
        "Est. cost is the AI cost of the run that produced the report, at the configured list prices.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
