"""The PQI results report: one row per candidate, for HR.

Who the candidate is (employee code, name, designation, location), the
headline result (Overall, Technical, Business, AFI, readiness), the ten
competency scores, flags and completion, and the report's top strengths and
development priorities — one sheet that can be filtered, sorted and shared.
Candidates who have not finished are listed on their own sheet.

The validation workbook (pqi_validation.py) is the technical counterpart for
checking the scoring engine; this one is for reading results. Read-only: no
AI calls. Designation comes from the employee master at the moment of
download, best effort — it is not stored with an attempt.
"""
from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

IST = timezone(timedelta(hours=5, minutes=30))
NAVY = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Same bands the readiness decision uses under Rubric v2 (HR pilot bands on C1),
# so a cell's colour and the candidate's readiness read the same way.
SCORE_FILLS = [
    ("lessThan", [50], "F8CBAD"),            # Not Yet Ready range
    ("between", [50, 64.999], "FFE699"),     # Developing
    ("between", [65, 79.999], "C6E0B4"),     # Ready
    ("greaterThanOrEqual", [80], "8FBC8F"),  # High Readiness
]
READINESS_FILLS = {
    "High Readiness": "8FBC8F", "Ready": "C6E0B4", "Developing": "FFE699", "Not Yet Ready": "F8CBAD",
}
STATUS_LABELS = {"in_progress": "Taking the test", "processing": "Being scored", "failed": "Scoring failed",
                 "completed": "Completed (no report)"}


def _rubric_label(report: dict) -> str:
    rubric = report.get("scoring_rubric") or "SRT_Scoring_Rubric_v1"
    return "v2" if rubric.endswith("v2") else "v1"


def _issued(report: dict) -> str:
    raw = report.get("generated_at") or ""
    try:
        dt = datetime.fromisoformat(raw.rstrip("Z")).replace(tzinfo=timezone.utc).astimezone(IST)
        return dt.strftime("%d-%b-%Y %H:%M")
    except ValueError:
        return raw


def _titles(items) -> str:
    return "; ".join(str(i.get("title", "")).strip() for i in items or [] if str(i.get("title", "")).strip())


def build_results(sessions: list, identities: dict, designations: dict, master: dict | None) -> bytes:
    """sessions: full PQI session dicts (with session_id); identities: database.session_identities;
    designations: {casefolded employee code: designation}; master: the current PQI master (competency names)."""
    competencies = [(c["code"], c["name"]) for c in (master or {}).get("competencies", [])] or \
                   [(f"C{i}", "") for i in range(1, 11)]

    done, pending = [], []
    for s in sessions:
        report = s.get("report") or {}
        (done if s.get("status") == "completed" and report.get("srt_results") else pending).append(s)
    done.sort(key=lambda s: -((s["report"].get("overall_pqi_score")) or 0))

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    generated = datetime.now(IST).strftime("%d-%b-%Y %H:%M IST")
    ws["A1"] = "RDC – PQI Competency Assessment Results"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = (f"Confidential — HR and authorised assessors only. Generated {generated}. {len(done)} completed, "
                f"{len(pending)} not completed (see 'Not completed'). Scores /100 on scale C1; see 'How to read'.")
    ws["A2"].font = Font(italic=True, color="595959")

    header = (["#", "Emp Code", "Name", "Designation", "Location", "Candidate type", "E-mail", "Assessment date",
               "Overall /100", "Technical /100", "Business /100", "AFI /100", "AFI interpretation", "Readiness",
               "Manual review", "Critical flags", "Answered"]
              + [f"{code} {name}".strip() for code, name in competencies]
              + ["Top strengths", "Development priorities", "Scoring rubric", "Report issued"])
    HEAD = 4
    ws.append([])
    ws.append(header)
    for cell in ws[HEAD]:
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[HEAD].height = 48

    for n, s in enumerate(done, start=1):
        r, cand = s["report"], s.get("candidate") or {}
        who = identities.get(s["session_id"]) or {}
        code = (who.get("employee_code") or "").strip()
        done_info = r.get("completion") or {}
        presented = done_info.get("questions_presented") or len(r.get("srt_results") or [])
        answered = done_info.get("questions_answered",
                                 sum(1 for x in r.get("srt_results") or [] if x.get("status") == "scored"))
        by_code = {c["code"]: c.get("score") for c in r.get("competency_scores") or []}
        ws.append([
            n, code or "—", cand.get("candidate_name"), designations.get(code.casefold(), "") if code else "",
            cand.get("plant_location"), "Employee" if (cand.get("candidate_type") or "employee") == "employee"
            else "External", who.get("captured_email") or "", cand.get("assessment_date"),
            r.get("overall_pqi_score"), r.get("technical_acumen"), r.get("business_acumen"),
            r.get("anti_firefighting_index"), r.get("afi_band") or "", r.get("overall_readiness"),
            "Yes" if (r.get("readiness") or {}).get("manual_review_required") else "",
            len(r.get("critical_flags") or []) or "", f"{answered}/{presented}",
            *[by_code.get(c) for c, _ in competencies],
            _titles(r.get("top_strengths")), _titles(r.get("development_priorities")),
            _rubric_label(r), _issued(r),
        ])

    last = HEAD + len(done)
    first_score = header.index("Overall /100") + 1
    comp_first = header.index("Answered") + 2
    comp_last = comp_first + len(competencies) - 1
    readiness_col = header.index("Readiness") + 1
    for row in ws.iter_rows(min_row=HEAD + 1, max_row=last):
        for cell in row:
            cell.border = BOX
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column > comp_last)
        for cell in row[first_score - 1:first_score + 3] + row[comp_first - 1:comp_last]:
            cell.number_format = "0.0"
            cell.alignment = Alignment(horizontal="center", vertical="top")
        band = row[readiness_col - 1]
        if band.value in READINESS_FILLS:
            band.fill = PatternFill("solid", fgColor=READINESS_FILLS[band.value])
            band.font = Font(bold=True)
    if done:
        ranges = [f"{get_column_letter(first_score)}{HEAD + 1}:{get_column_letter(first_score + 2)}{last}",
                  f"{get_column_letter(comp_first)}{HEAD + 1}:{get_column_letter(comp_last)}{last}"]
        for rng in ranges:            # AFI has its own bands, so it is left uncoloured
            for operator, values, colour in SCORE_FILLS:
                ws.conditional_formatting.add(rng, CellIsRule(operator=operator, formula=[str(v) for v in values],
                                                              fill=PatternFill("solid", fgColor=colour)))
    ws.auto_filter.ref = f"A{HEAD}:{get_column_letter(len(header))}{max(last, HEAD + 1)}"
    ws.freeze_panes = ws.cell(row=HEAD + 1, column=4)
    widths = {"#": 4, "Emp Code": 11, "Name": 24, "Designation": 24, "Location": 18, "Candidate type": 11,
              "E-mail": 26, "Assessment date": 12, "AFI interpretation": 20, "Readiness": 15, "Manual review": 9,
              "Critical flags": 9, "Answered": 9, "Top strengths": 50, "Development priorities": 50,
              "Scoring rubric": 9, "Report issued": 16}
    for i, h in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 12 if i >= comp_first else 10)

    # ── Not completed ────────────────────────────────────────────────────────
    nc = wb.create_sheet("Not completed")
    nc.append(["Emp Code", "Name", "Location", "Candidate type", "Status", "Answers saved", "Started", "Note"])
    for cell in nc[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for s in sorted(pending, key=lambda s: s.get("status") or ""):
        cand, who = s.get("candidate") or {}, identities.get(s["session_id"]) or {}
        saved = sum(1 for v in (s.get("collected_answers") or {}).values() if str(v or "").strip())
        nc.append([who.get("employee_code") or "—", cand.get("candidate_name"), cand.get("plant_location"),
                   "Employee" if (cand.get("candidate_type") or "employee") == "employee" else "External",
                   STATUS_LABELS.get(s.get("status"), s.get("status")), f"{saved}/{len(s.get('questions') or [])}",
                   cand.get("assessment_date"), (s.get("error") or "")[:200]])
    for col, width in zip("ABCDEFGH", (11, 24, 18, 11, 18, 12, 12, 70)):
        nc.column_dimensions[col].width = width

    # ── Competency key ───────────────────────────────────────────────────────
    key = wb.create_sheet("Competency key")
    key.append(["Code", "Competency", "Lens", "Acumen weight", "Definition"])
    for cell in key[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for c in (master or {}).get("competencies", []):
        key.append([c["code"], c["name"], c["lens"], c["weight"], c.get("definition", "")])
    for col, width in zip("ABCDE", (7, 50, 11, 13, 110)):
        key.column_dimensions[col].width = width

    # ── How to read ──────────────────────────────────────────────────────────
    notes = wb.create_sheet("How to read")
    afi_bands = "; ".join(f"{b['lower']}–{b['upper']} {b['label']}" for b in (master or {}).get("afi_bands", []))
    for line in [
        "How to read the PQI results",
        "",
        "Each of the 30 situations is scored 0–9 and converted on scale C1 (0, 20, 35, 50, 60, 70, 80, 88, 94, 100). "
        "A competency score is the average of its three situations; Overall is the average of the ten competencies, "
        "equally weighted.",
        "Technical = C6–C10 (C7, C9, C10 weighted 1.5). Business = C1–C5 (C2, C3, C5 weighted 1.5).",
        "Readiness (Rubric v2, pilot bands): 80+ High Readiness · 65–79 Ready · 50–64 Developing · below 50 Not Yet "
        "Ready. Bands are pilot-calibrated. Score cells are coloured on the same bands.",
        f"AFI (Anti-Firefighting Index) is reported separately and is not part of Overall. Interpretation: {afi_bands}.",
        "Unanswered situations count as zero. 'Answered' shows how many of the 30 were answered.",
        "Manual review = the report asks a person to read it before relying on readiness (for example an unconfirmed "
        "Critical Failure). Critical flags = confirmed Critical Failures.",
        "Scoring rubric: v2 is HR's revised engine (October 2026). A v1 row was scored by the earlier engine and is "
        "not directly comparable — re-score it (Full) to bring it onto v2.",
        "Designation is the employee master's current designation at the time of download.",
        "Confidential — for HR and authorised assessors only.",
    ]:
        notes.append([line])
    notes["A1"].font = Font(bold=True, size=13, color=NAVY)
    notes.column_dimensions["A"].width = 150
    for row in notes.iter_rows(min_row=3):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
