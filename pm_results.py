"""The Plant Manager results report: one row per candidate, for HR.

The same layout as the PQI results report (pqi_results.py) so the two read
alike: who the candidate is, the headline result (total /300, score /100,
readiness), the ten competency scores, flags and completion, strengths and
development areas, plus sheets for candidates not yet completed, the
competency key and how to read the scores.

Read-only — it reports what the Plant Manager pipeline already stored and
changes nothing about how Plant Manager is scored. Designation comes from the
employee master at the moment of download, best effort.
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pqi_results import BOX, HEADER_FILL, HEADER_FONT, IST, NAVY, not_completed_sheet

# The Plant Manager tiers, top down, as main.READINESS_TIERS applies them.
READINESS_FILLS = {
    "Ready for Higher Responsibility": "8FBC8F",
    "Ready to be Plant Manager":       "C6E0B4",
    "Ready with Structured Support":   "FFE699",
    "Not Yet Ready":                   "F8CBAD",
    "Low Potential":                   "E6A0A0",
}
# Competency /10, on the tiers' competency floors (5.0 / 6.0 / 6.5).
COMPETENCY_FILLS = [
    ("lessThan", [5], "F8CBAD"),
    ("between", [5, 5.999], "FFE699"),
    ("between", [6, 6.499], "C6E0B4"),
    ("greaterThanOrEqual", [6.5], "8FBC8F"),
]
# Total /300, on the tiers' total bands (90 / 150 / 210 / 251).
TOTAL_FILLS = [
    ("lessThan", [90], "E6A0A0"),
    ("between", [90, 149.999], "F8CBAD"),
    ("between", [150, 209.999], "FFE699"),
    ("between", [210, 250.999], "C6E0B4"),
    ("greaterThanOrEqual", [251], "8FBC8F"),
]


def _items(items, key, extra=None) -> str:
    out = []
    for i in items or []:
        if not isinstance(i, dict):          # some reports hold plain strings here
            i = {key: i}
        text = str(i.get(key) or "").strip()
        if text:
            if extra and i.get(extra):
                text += f" ({i[extra]})"
            out.append(text)
    return "; ".join(out)


def build_results(sessions: list, identities: dict, designations: dict, competencies: list, tiers: list) -> bytes:
    """sessions: full Plant Manager session dicts (with session_id); identities: database.session_identities;
    designations: {casefolded employee code: designation}; competencies: the ten names in report order;
    tiers: main.READINESS_TIERS — (name, minimum total /300, competency floor), top down."""
    done, pending = [], []
    for s in sessions:
        report = s.get("report") or {}
        (done if s.get("status") == "completed" and report.get("overall_score_out_of_300") is not None
         else pending).append(s)
    done.sort(key=lambda s: -(s["report"].get("overall_score_out_of_300") or 0))

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    generated = datetime.now(IST).strftime("%d-%b-%Y %H:%M IST")
    ws["A1"] = "RDC – Plant Manager Competency Assessment Results (PLANT-PRO)"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = (f"Confidential — HR and authorised assessors only. Generated {generated}. {len(done)} completed, "
                f"{len(pending)} not completed (see 'Not completed'). Total /300 is English-adjusted; see "
                "'How to read'.")
    ws["A2"].font = Font(italic=True, color="595959")

    header = (["#", "Emp Code", "Name", "Designation", "Location", "Candidate type", "E-mail", "Assessment date",
               "Total /300", "Score /100", "Base total /300 (before English)", "Readiness", "Readiness note",
               "Integrity flag", "Weakest competency", "Answered"]
              + list(competencies)
              + ["Top strengths", "Development areas"])
    HEAD = 4
    ws.append([])
    ws.append(header)
    for cell in ws[HEAD]:
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[HEAD].height = 60

    for n, s in enumerate(done, start=1):
        r, cand = s["report"], s.get("candidate") or {}
        who = identities.get(s["session_id"]) or {}
        code = (who.get("employee_code") or "").strip()
        answered = sum(1 for v in (s.get("collected_answers") or {}).values() if str(v or "").strip())
        flag = r.get("integrity_red_flag") or {}
        summary = r.get("competency_summary") or {}
        note = r.get("readiness_explanation") or ""
        if r.get("readiness_demoted_from"):
            note = f"Lowered from {r['readiness_demoted_from']}: {note}"
        ws.append([
            n, code or "—", cand.get("candidate_name"), designations.get(code.casefold(), "") if code else "",
            cand.get("plant_location"), "Employee" if (cand.get("candidate_type") or "employee") == "employee"
            else "External", who.get("captured_email") or "", cand.get("assessment_date"),
            r.get("overall_score_out_of_300"), r.get("normalized_score_out_of_100"), r.get("base_score_out_of_300"),
            r.get("overall_readiness"), note,
            ("Yes — " + str(flag.get("evidence") or "").strip()).rstrip(" —") if flag.get("present") else "",
            r.get("weakest_competency") or "", f"{answered}/{len(s.get('questions') or []) or 30}",
            *[summary.get(c) for c in competencies],
            _items(r.get("top_strengths"), "strength"), _items(r.get("development_areas"), "area", "priority"),
        ])

    last = HEAD + len(done)
    total_col = header.index("Total /300") + 1
    comp_first = header.index("Answered") + 2
    comp_last = comp_first + len(competencies) - 1
    readiness_col = header.index("Readiness") + 1
    wrap_cols = {header.index(h) + 1 for h in ("Readiness note", "Integrity flag", "Top strengths",
                                                "Development areas")}
    for row in ws.iter_rows(min_row=HEAD + 1, max_row=last):
        for cell in row:
            cell.border = BOX
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_cols)
        for cell in row[total_col - 1:total_col + 2] + row[comp_first - 1:comp_last]:
            cell.number_format = "0.0"
            cell.alignment = Alignment(horizontal="center", vertical="top")
        band = row[readiness_col - 1]
        if band.value in READINESS_FILLS:
            band.fill = PatternFill("solid", fgColor=READINESS_FILLS[band.value])
            band.font = Font(bold=True)
    if done:
        for rng, fills in (
            (f"{get_column_letter(total_col)}{HEAD + 1}:{get_column_letter(total_col)}{last}", TOTAL_FILLS),
            (f"{get_column_letter(comp_first)}{HEAD + 1}:{get_column_letter(comp_last)}{last}", COMPETENCY_FILLS),
        ):
            for operator, values, colour in fills:
                ws.conditional_formatting.add(rng, CellIsRule(operator=operator, formula=[str(v) for v in values],
                                                              fill=PatternFill("solid", fgColor=colour)))
    ws.auto_filter.ref = f"A{HEAD}:{get_column_letter(len(header))}{max(last, HEAD + 1)}"
    ws.freeze_panes = ws.cell(row=HEAD + 1, column=4)
    widths = {"#": 4, "Emp Code": 11, "Name": 24, "Designation": 24, "Location": 18, "Candidate type": 11,
              "E-mail": 26, "Assessment date": 12, "Readiness": 17, "Readiness note": 40, "Integrity flag": 24,
              "Weakest competency": 22, "Answered": 9, "Top strengths": 50, "Development areas": 50}
    for i, h in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 13 if i >= comp_first else 11)

    not_completed_sheet(wb, pending, identities)

    key = wb.create_sheet("Competency key")
    key.append(["#", "Competency"])
    for cell in key[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for i, name in enumerate(competencies, start=1):
        key.append([i, name])
    key.column_dimensions["A"].width, key.column_dimensions["B"].width = 5, 50

    notes = wb.create_sheet("How to read")
    tier_lines = []
    for name, minimum, floor in tiers:
        if minimum <= 0:
            tier_lines.append(f"  {name}: below the next band")
        else:
            tier_lines.append(f"  {name}: total {minimum}+ /300"
                              + (f" and every competency at least {floor:g}/10" if floor else ""))
    for line in [
        "How to read the Plant Manager results",
        "",
        "Each of the 30 situations is scored out of 10. Total /300 is the sum after the English-proficiency "
        "adjustment (15% weight); Score /100 is the same total as a percentage. 'Base total' is the content score "
        "before that adjustment.",
        "A competency score is the average of its three situations, out of 10.",
        "Readiness, top down — the first tier whose total AND competency floor are both met:",
        *tier_lines,
        "'Readiness note' says when a competency below its floor lowered the tier.",
        "Integrity flag = an explicit severe integrity red flag in the answers; it caps readiness at "
        "'Ready with Structured Support' or below.",
        "Unanswered situations count as zero. 'Answered' shows how many of the 30 were answered.",
        "Colours: Total on the tier bands (90 / 150 / 210 / 251); competencies on the floors (5.0 / 6.0 / 6.5).",
        "Designation is the employee master's current designation at the time of download.",
        "Confidential — for HR and authorised assessors only.",
    ]:
        notes.append([line])
    notes["A1"].font = Font(bold=True, size=13, color=NAVY)
    notes.column_dimensions["A"].width = 150
    for row in notes.iter_rows(min_row=3):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
